import io
import json
from datetime import date, timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import User, Worker
from production.daily_dazmol import compute_daily_dazmol, get_day_window, get_patok_accepted_pastals
from production.daily_image_reports import get_branch_daily_report_data
from production.models import (
    Article, ArticleOperation, Box, CuttingBatch, CuttingBatchItem, DailyPatokProgress, Operation,
    OperationGroup, Order, OrderItem, OrderItemSize, Ticket,
)


class DailyDazmolRulesTest(TestCase):
    """
    Asosiy dazmol = patokka qabul qilingan pastallar dazmoli (qayerda bo'lsa ham): bot yuborgan pastal planshetda
    qabul qilinganda (Koson ham, Uychi ham) yoki Uychi omboridan U-patokka topshirilganda. Patok dazmoli = terminal
    bo'yicha. Egasiz = hech bir patokka qabul qilinmagan pastallar dazmoli.
    /sewing-statistics/daily/ dagi asosiy dazmol /daily-report/ dagi shablonlar yig'indisiga teng bo'lishi shart.
    """

    def setUp(self):
        self.admin = User.objects.create_superuser(username='dz_admin', password='x', role=User.Role.SUPER_ADMIN)
        self.client.force_login(self.admin)
        self.worker = Worker.objects.create(worker_id='W-DZ-1', first_name='Dazmol', last_name='Chi')

        self.group_f = OperationGroup.objects.create(name='FUTBOLKA KATTA')
        self.group_t = OperationGroup.objects.create(name='TUNIKA')
        self.art_f = Article.objects.create(code='ART-F', name='Futbolka')
        self.art_t = Article.objects.create(code='ART-T', name='Tunika')
        # save() xukisiz: shablon sinxronlashi operatsiyalarni o'zgartirmasin
        Article.objects.filter(pk=self.art_f.pk).update(operation_group=self.group_f)
        Article.objects.filter(pk=self.art_t.pk).update(operation_group=self.group_t)
        self.art_f.refresh_from_db()
        self.art_t.refresh_from_db()

        op_dazmol = Operation.objects.create(code='OP-DZ', name='DAZMOL')
        op_tasma = Operation.objects.create(code='OP-TS', name='TASMA DAZMOL')
        self.dz_f = ArticleOperation.objects.create(article=self.art_f, operation=op_dazmol, price_per_unit=Decimal('100'), sequence=9)
        self.dz_t = ArticleOperation.objects.create(article=self.art_t, operation=op_dazmol, price_per_unit=Decimal('100'), sequence=9)
        self.tasma_t = ArticleOperation.objects.create(article=self.art_t, operation=op_tasma, price_per_unit=Decimal('50'), sequence=5)

        self.order = Order.objects.create(order_number='ORD-DZ', article=self.art_f, total_quantity=1000)
        self.item_f = OrderItem.objects.create(order=self.order, article=self.art_f, quantity=500)
        self.item_t = OrderItem.objects.create(order=self.order, article=self.art_t, quantity=500)
        self.size_f = OrderItemSize.objects.create(order_item=self.item_f, size_name='M', planned_quantity=500)
        self.size_t = OrderItemSize.objects.create(order_item=self.item_t, size_name='M', planned_quantity=500)
        self._box_no = 0
        self._batch_no = 0

        now = timezone.now()
        # P-A: Futbolka, K1 planshetda qabul qilgan; K2 va K5 da dazmollandi
        a_boxes = self._pastal('P-A', self.item_f, self.size_f, screen=1, status='ACCEPTED', received=True, qtys=[50, 50, 20], branch='HQ')
        self._scan(a_boxes[0], self.dz_f, 2, now)
        self._scan(a_boxes[1], self.dz_f, 5, now)
        # P-B: Tunika, K1 qabul qilgan; K1 da dazmol + tasma dazmol (hisobga olinmaydi)
        b_boxes = self._pastal('P-B', self.item_t, self.size_t, screen=1, status='ACCEPTED', received=True, qtys=[30], branch='HQ')
        self._scan(b_boxes[0], self.dz_t, 1, now)
        self._scan(b_boxes[0], self.tasma_t, 1, now)
        # P-C: K3 ga yuborilgan, lekin qabul qilinmagan -> egasiz
        c_boxes = self._pastal('P-C', self.item_f, self.size_f, screen=3, status='DISPATCHED', received=False, qtys=[40], branch='HQ')
        self._scan(c_boxes[0], self.dz_f, 3, now)
        # P-D: Uychi ombori yuk xatini qabul qilib U1 ga topshirgan (planshet qabuli yo'q) -> U1 asosiy
        d_boxes = self._pastal('P-D', self.item_f, self.size_f, screen=14, status='ACCEPTED', received=False, qtys=[70], branch='UYCHI')
        self._scan(d_boxes[0], self.dz_f, 14, now)
        # P-W: Uychi omborida, hali hech bir patokka topshirilmagan -> egasiz
        w_boxes = self._pastal('P-W', self.item_f, self.size_f, screen=None, status='ACCEPTED', received=False, qtys=[25], branch='UYCHI')
        self._scan(w_boxes[0], self.dz_f, 15, now)
        # P-K: Koson pastali ACCEPTED, lekin qutilari planshetda qabul qilinmagan -> egasiz (Uychi istisnosi Kosonga o'tmaydi)
        k_boxes = self._pastal('P-K', self.item_f, self.size_f, screen=2, status='ACCEPTED', received=False, qtys=[15], branch='HQ')
        self._scan(k_boxes[0], self.dz_f, 2, now)
        # Kecha 23:30 dagi dazmol (P-A uchinchi qutisi) bugunga kirmasligi kerak
        yesterday = timezone.localdate() - timedelta(days=1)
        late_yesterday = get_day_window(yesterday)[1] - timedelta(minutes=30)
        self._scan(a_boxes[2], self.dz_f, 1, late_yesterday)

        self.batch_a = a_boxes[0].cutting_batch_item.batch
        self.b_boxes = b_boxes

    # --- yordamchilar -------------------------------------------------------------------------------
    def _pastal(self, code, order_item, size, screen, status, received, qtys, branch):
        self._batch_no += 1
        batch = CuttingBatch.objects.create(
            order_item=order_item, batch_number=self._batch_no, pastal_code=code,
            assigned_screen=screen, assigned_branch=branch, dispatch_status=status,
        )
        item = CuttingBatchItem.objects.create(batch=batch, order_item_size=size, quantity=sum(qtys), real_quantity=sum(qtys))
        boxes = []
        for qty in qtys:
            self._box_no += 1
            boxes.append(Box.objects.create(
                order=self.order, article=order_item.article, cutting_batch_item=item, box_number=self._box_no,
                quantity=qty, pastal_number=code, assigned_screen=screen, is_received_by_patok=received,
            ))
        return boxes

    def _scan(self, box, ao, screen, when):
        return Ticket.objects.create(
            box=box, article_operation=ao, quantity=box.quantity, worker=self.worker,
            status=Ticket.Status.SCANNED, screen_number=screen, scanned_at=when,
        )

    def _report_rows(self, day=None):
        day = day or timezone.localdate()
        rows = {}
        for branch in ('HQ', 'UYCHI'):
            for r in get_branch_daily_report_data(branch, day)['rows']:
                rows[r['screen_number']] = r
        return rows

    # --- testlar ------------------------------------------------------------------------------------
    def test_compute_splits_asosiy_patok_and_egasiz(self):
        dz = compute_daily_dazmol(*get_day_window(timezone.localdate()))

        self.assertEqual(dz.asosiy_by_patok, {1: 130, 14: 70})
        self.assertEqual([(t['name'], t['qty']) for t in dz.templates(1)], [('FUTBOLKA KATTA', 100), ('TUNIKA', 30)])
        self.assertEqual([(t['name'], t['qty']) for t in dz.templates(14)], [('FUTBOLKA KATTA', 70)])
        self.assertEqual(dz.pastals(1), [{'code': 'P-A', 'qty': 100}, {'code': 'P-B', 'qty': 30}])
        self.assertEqual(dz.pastals(14), [{'code': 'P-D', 'qty': 70}])
        self.assertEqual(dz.patok_dazmol_by_screen, {1: 30, 2: 65, 3: 40, 5: 50, 14: 70, 15: 25})
        self.assertEqual(dz.egasiz_by_screen, {2: 15, 3: 40, 15: 25})
        self.assertEqual(sum(dz.asosiy_by_patok.values()) + sum(dz.egasiz_by_screen.values()),
                         sum(dz.patok_dazmol_by_screen.values()))

    def test_daily_page_asosiy_equals_report_template_sum_for_every_patok(self):
        res = self.client.get(reverse('production:sewing_statistics_daily') + '?all=1')
        self.assertEqual(res.status_code, 200)
        report_rows = self._report_rows()

        for p in res.context['patoks_data']:
            row = report_rows.get(p['screen_number'])
            template_sum = sum(prod['dazmol_qty'] for prod in row['products']) if row else 0
            self.assertEqual(p['primary_dazmol_total'], template_sum, f"{p['patok_code']} mos emas")
            self.assertEqual(sum(i['qty'] for i in p['pastal_dazmol_items']), p['primary_dazmol_total'])

        by_sn = {p['screen_number']: p for p in res.context['patoks_data']}
        self.assertEqual(by_sn[1]['primary_dazmol_total'], 130)
        self.assertEqual(by_sn[2]['primary_dazmol_total'], 0)
        self.assertEqual(by_sn[5]['primary_dazmol_total'], 0)
        self.assertEqual(by_sn[14]['primary_dazmol_total'], 70)
        self.assertEqual(by_sn[2]['patok_dazmol_qty'], 65)
        # Qabul qilinmagan pastal asosiy emas, faqat qo'shimcha
        self.assertNotIn('P-C', [x['pastal_code'] for x in by_sn[3]['primary_pastals']])
        self.assertIn('P-C', [x['pastal_code'] for x in by_sn[3]['extra_pastals']])
        self.assertNotIn('P-K', [x['pastal_code'] for x in by_sn[2]['primary_pastals']])
        self.assertIn('P-D', [x['pastal_code'] for x in by_sn[14]['primary_pastals']])
        # Ikki marta sanash yo'q
        self.assertEqual(res.context['kpis']['primary_dazmol_qty'], 200)
        self.assertEqual(res.context['kpis']['egasiz_dazmol_qty'], 80)

    def test_daily_page_template_lines_equal_report_lines(self):
        res = self.client.get(reverse('production:sewing_statistics_daily') + '?all=1')
        report_rows = self._report_rows()

        checked = 0
        for p in res.context['patoks_data']:
            if not p['primary_dazmol_total']:
                self.assertEqual(p['dazmol_templates'], [])
                continue
            row = report_rows[p['screen_number']]
            self.assertEqual([(t['name'], t['qty']) for t in p['dazmol_templates']],
                             [(x['shablon_name'], x['dazmol_qty']) for x in row['products']])
            self.assertEqual(sum(t['qty'] for t in p['dazmol_templates']), p['primary_dazmol_total'])
            checked += 1
        self.assertEqual(checked, 2)   # K1 va U1

    def test_patok_filter_is_exact_and_shows_idle_patok(self):
        url = reverse('production:sewing_statistics_daily')
        res = self.client.get(url + '?patok=K1')
        self.assertEqual([p['patok_code'] for p in res.context['patoks_data']], ['K1'])   # K10..K13 emas
        res = self.client.get(url + '?patok=K10')                                       # bugun ishlamagan
        self.assertEqual([p['patok_code'] for p in res.context['patoks_data']], ['K10'])

    def test_report_page_links_to_patok_and_shows_factory_total(self):
        res = self.client.get(reverse('production:sewing_statistics_daily_report'))
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['factory']['total_dazmol'], 200)
        self.assertContains(res, '&patok=K1"')
        res = self.client.get(reverse('production:sewing_statistics_daily_report') + '?tab=uychi')
        self.assertContains(res, '&patok=U1"')

    def test_report_rows_show_asosiy_by_template(self):
        rows = self._report_rows()
        self.assertEqual(rows[1]['dazmol_total'], 130)
        self.assertEqual([(p['shablon_name'], p['dazmol_qty']) for p in rows[1]['products']],
                         [('FUTBOLKA KATTA', 100), ('TUNIKA', 30)])
        self.assertEqual(rows[3]['dazmol_total'], 0)     # yuborilgan, qabul qilinmagan
        self.assertEqual(rows[2]['dazmol_total'], 0)     # Koson: planshet qabuli yo'q
        self.assertEqual(rows[14]['dazmol_total'], 70)   # Uychi ombori U1 ga topshirgan
        self.assertEqual([(p['shablon_name'], p['dazmol_qty']) for p in rows[14]['products']], [('FUTBOLKA KATTA', 70)])
        self.assertEqual(rows[15]['dazmol_total'], 0)    # hali Uychi omborida

    def test_rejected_or_redispatched_pastal_leaves_asosiy(self):
        Box.objects.filter(pk__in=[b.pk for b in self.b_boxes]).update(is_received_by_patok=False)   # P-B qaytarildi
        self.assertEqual(compute_daily_dazmol(*get_day_window(timezone.localdate())).asosiy_by_patok, {1: 100, 14: 70})

        CuttingBatch.objects.filter(pk=self.batch_a.pk).update(dispatch_status=CuttingBatch.DispatchStatus.DISPATCHED)
        self.assertEqual(compute_daily_dazmol(*get_day_window(timezone.localdate())).asosiy_by_patok, {14: 70})

        CuttingBatch.objects.filter(pastal_code='P-D').update(dispatch_status=CuttingBatch.DispatchStatus.REJECTED)
        self.assertEqual(compute_daily_dazmol(*get_day_window(timezone.localdate())).asosiy_by_patok, {})

    def test_uychi_bot_dispatch_shows_on_tablet_and_counts_after_accept(self):
        """Telegram bot U17 (30-ekran) ga yuborgan pastal planshetda kutilmoqda; master qabul qilgach asosiyga kiradi."""
        boxes = self._pastal('P-U17', self.item_f, self.size_f, screen=30, status='DISPATCHED', received=False, qtys=[45], branch='UYCHI')
        self._scan(boxes[0], self.dz_f, 30, timezone.now())
        window = get_day_window(timezone.localdate())
        self.assertEqual((compute_daily_dazmol(*window).asosiy(30), compute_daily_dazmol(*window).egasiz(30)), (0, 45))

        detail = self.client.get(reverse('production:master_tablet_patok_detail', args=[30])).json()
        self.assertIn('P-U17', [w['pastal_code'] for w in detail['waiting_pastals']])

        res = self.client.post(reverse('production:master_tablet_accept'), content_type='application/json',
                               data=json.dumps({'batch_id': boxes[0].cutting_batch_item.batch_id, 'screen_number': 30}))
        self.assertTrue(res.json()['success'])

        dz = compute_daily_dazmol(*window)
        self.assertEqual((dz.asosiy(30), dz.egasiz(30)), (45, 0))
        self.assertEqual(self._report_rows()[30]['dazmol_total'], 45)

    def test_uychi_handout_counts_as_acceptance_only_in_uychi(self):
        dz = compute_daily_dazmol(*get_day_window(timezone.localdate()))
        self.assertEqual(dz.asosiy(14), 70)   # P-D: ombor U1 ga topshirgan
        self.assertEqual(dz.egasiz(15), 25)   # P-W: hali omborda
        self.assertEqual(dz.egasiz(2), 15)    # P-K: Koson, planshet qabuli yo'q
        # Queryset qoidasi (1-sahifadagi asosiy pastallar) Python qoidasi bilan bir xil
        accepted = get_patok_accepted_pastals()
        self.assertEqual(accepted.get(14), {'P-D'})
        self.assertNotIn('P-K', accepted.get(2, set()))
        self.assertNotIn('P-W', set().union(*accepted.values()))

    def test_only_exact_dazmol_counts_planka_is_not_finished_work(self):
        """PLANKA DAZMOL — oraliq amal: asosiy ham, patok dazmoli ham emas; model sahifasida ham yakuniy dazmol emas."""
        from production.sewing_statistics_views import get_dazmol_operation_ids_for_article
        op_planka = Operation.objects.create(code='OP-PL', name='PLANKA DAZMOL')
        planka_f = ArticleOperation.objects.create(article=self.art_f, operation=op_planka, price_per_unit=Decimal('20'), sequence=4)
        boxes = self._pastal('P-PL', self.item_f, self.size_f, screen=1, status='ACCEPTED', received=True, qtys=[10, 10], branch='HQ')
        now = timezone.now()
        self._scan(boxes[0], planka_f, 1, now)      # faqat planka — hali tayyor emas
        self._scan(boxes[1], planka_f, 1, now)      # planka + yakuniy dazmol — bir marta sanaladi
        self._scan(boxes[1], self.dz_f, 1, now)

        dz = compute_daily_dazmol(*get_day_window(timezone.localdate()))
        self.assertEqual(dz.asosiy(1), 130 + 10)
        self.assertEqual(dz.patok_dazmol(1), 30 + 10)
        self.assertEqual(get_dazmol_operation_ids_for_article(self.art_f), {self.dz_f.id})

    def test_day_window_keeps_yesterday_separate(self):
        yesterday = timezone.localdate() - timedelta(days=1)
        self.assertEqual(compute_daily_dazmol(*get_day_window(yesterday)).asosiy_by_patok, {1: 20})

    def test_both_pages_store_same_dazmol_for_monthly_excel(self):
        today = timezone.localdate()
        self.client.get(reverse('production:sewing_statistics_daily'))
        self.assertEqual(DailyPatokProgress.objects.get(screen_number=1, date=today).dazmol_qty, 130)
        self.client.get(reverse('production:sewing_statistics_daily_report'))
        self.assertEqual(DailyPatokProgress.objects.get(screen_number=1, date=today).dazmol_qty, 130)
        self.assertEqual(DailyPatokProgress.objects.get(screen_number=14, date=today).dazmol_qty, 70)
        self.client.get(reverse('production:sewing_statistics_daily'))
        self.assertEqual(DailyPatokProgress.objects.get(screen_number=1, date=today).dazmol_qty, 130)


class MonthlyExcelUychiRowsTest(TestCase):
    """Oylik Excel: Uychi'da U1..U21 doim, undan keyingi patok shu oyda ishlagan bo'lsa qatori qo'shiladi."""

    def test_patok_after_u21_gets_a_row_and_rows_match_total(self):
        import openpyxl
        from production.excel_reports import generate_monthly_patok_production_excel

        day = date(2026, 10, 5)
        for sn, qty in ((14, 100), (36, 40)):   # U1 va U23
            DailyPatokProgress.objects.create(screen_number=sn, date=day, patok_code=f'U{sn - 13}', dazmol_qty=qty, davomat=5)

        content = generate_monthly_patok_production_excel(2026, 10)
        ws = openpyxl.load_workbook(io.BytesIO(content.getvalue()))['Уйчи тумани']
        labels = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
        self.assertIn('23-паток', labels)
        self.assertNotIn('24-паток', labels)

        total_row = 4 + labels.index('ЖАМИ:')
        day_col = 1 + day.day
        rows_sum = sum(ws.cell(row=r, column=day_col).value or 0 for r in range(4, total_row))
        self.assertEqual(rows_sum, ws.cell(row=total_row, column=day_col).value)
        self.assertEqual(rows_sum, 140)
