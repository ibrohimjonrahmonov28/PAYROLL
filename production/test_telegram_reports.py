import io
import datetime
import tempfile
from decimal import Decimal
from unittest.mock import patch, MagicMock

from django.test import TestCase
from django.utils import timezone
from django.core.management import call_command
from django.urls import reverse

import openpyxl

from accounts.models import User, Worker
from production.models import Customer, Order, Box, Article, Operation, ArticleOperation, Ticket
from production.excel_reports import generate_daily_excel_report
from production.telegram_reports import send_daily_excel_report


class TelegramDailyReportTest(TestCase):
    def setUp(self):
        # Admin user
        self.superadmin = User.objects.create_superuser(
            username='admin_test',
            password='password123',
            role=User.Role.SUPER_ADMIN
        )

        # Worker
        self.worker = Worker.objects.create(
            worker_id='W-999',
            first_name='Dilshod',
            last_name='Karimov',
            phone_number='+998901234567'
        )

        # Order & Box
        self.customer = Customer.objects.create(name='Test Customer')
        self.order = Order.objects.create(
            order_number='ORD-TEST-001',
            customer=self.customer
        )
        self.box = Box.objects.create(
            order=self.order,
            box_number=1,
            box_code='BX-001',
            quantity=100
        )

        # Article & Operation
        self.article = Article.objects.create(
            code='ART-99',
            name='Test Model Futbolka'
        )
        self.op = Operation.objects.create(
            code='OP-01',
            name='Yoqa tikish'
        )
        self.art_op = ArticleOperation.objects.create(
            article=self.article,
            operation=self.op,
            price_per_unit=Decimal('500.00'),
            sequence=1
        )

        # Scanned Ticket
        self.today = timezone.localdate()
        self.ticket = Ticket.objects.create(
            ticket_code='TK-TEST-001',
            stiker_code='STIK9999',
            box=self.box,
            article_operation=self.art_op,
            quantity=50,
            price_per_unit=Decimal('500.00'),
            total_amount=Decimal('25000.00'),
            status=Ticket.Status.SCANNED,
            worker=self.worker,
            scanned_at=timezone.now()
        )

    def test_generate_daily_excel_report(self):
        """Excel hisoboti to'g'ri shakllanadi va 2 ta varaqdan iborat bo'ladi"""
        buf = generate_daily_excel_report(target_date=self.today)
        self.assertIsInstance(buf, io.BytesIO)
        
        # openpyxl bilan o'qib tekshirish
        wb = openpyxl.load_workbook(buf)
        self.assertIn("Xodimlar Kunlik Hisoboti", wb.sheetnames)
        self.assertIn("Skanerlangan Stikerlar", wb.sheetnames)

        # 1-varaq tekshiruvi
        ws1 = wb["Xodimlar Kunlik Hisoboti"]
        found_worker = False
        found_sticker = False
        for row in ws1.iter_rows(values_only=True):
            if 'W-999' in row:
                found_worker = True
                # Stiker kodi qatorda bo'lishi kerak
                for cell_val in row:
                    if cell_val and 'STIK9999' in str(cell_val):
                        found_sticker = True
        self.assertTrue(found_worker, "Worker W-999 Excel 1-varaqda topilmadi")
        self.assertTrue(found_sticker, "Stiker kodi STIK9999 1-varaqda topilmadi")

        # 2-varaq tekshiruvi
        ws2 = wb["Skanerlangan Stikerlar"]
        found_ticket_detail = False
        for row in ws2.iter_rows(values_only=True):
            if 'STIK9999' in row:
                found_ticket_detail = True
                self.assertIn('Yoqa tikish', row)
                self.assertIn('Dilshod Karimov', row)
        self.assertTrue(found_ticket_detail, "Stiker tafsiloti 2-varaqda topilmadi")

    def test_send_daily_excel_report_no_credentials(self):
        """Token yoki chat_id berilmasa, xatolik xabari qaytishi kerak"""
        res = send_daily_excel_report(target_date=self.today, chat_id="", bot_token="")
        self.assertFalse(res['success'])
        self.assertIn("BOT_TOKEN sozlanmagan", res['message'])

        res2 = send_daily_excel_report(target_date=self.today, chat_id="", bot_token="12345:TEST_TOKEN")
        self.assertFalse(res2['success'])
        self.assertIn("TELEGRAM_REPORT_CHAT_ID sozlanmagan", res2['message'])

    @patch('httpx.Client.post')
    def test_send_daily_excel_report_success(self, mock_post):
        """Telegram API ga to'g'ri so'rov jo'natiladi"""
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.json.return_value = {'ok': True, 'result': {'message_id': 123}}
        mock_post.return_value = mock_response

        res = send_daily_excel_report(
            target_date=self.today,
            chat_id="-1001234567890",
            bot_token="12345:TEST_TOKEN"
        )
        self.assertTrue(res['success'])
        self.assertTrue(mock_post.called)

    def test_management_command(self):
        """Management command xatosiz ishlaydi va --save-local fayl yaratadi"""
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tf:
            tmp_path = tf.name

        call_command('send_daily_report', date=self.today.strftime('%Y-%m-%d'), save_local=tmp_path)
        
        # Fayl yaratilganligini tekshirish
        wb = openpyxl.load_workbook(tmp_path)
        self.assertIn("Xodimlar Kunlik Hisoboti", wb.sheetnames)

    def test_superadmin_views(self):
        """Superadmin panelidagi download va telegram jo'natish viewlari ishlaydi"""
        self.client.force_login(self.superadmin)

        # Download view
        res_dl = self.client.get(reverse('superadmin_download_daily_excel'))
        self.assertEqual(res_dl.status_code, 200)
        self.assertEqual(res_dl['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # Telegram report view (token yo'q bo'lsa ogohlantirish qaytadi)
        res_tg = self.client.get(reverse('superadmin_send_telegram_report'))
        self.assertEqual(res_tg.status_code, 302)

